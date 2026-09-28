from pathlib import Path
import datetime as dt
import hashlib
import json
import re
import sys
import unicodedata
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(r'E:\projects\knowledgeProjects\DailyWrite')
SOURCE = Path(r'D:\webDownload\作品列表导出 (1).xlsx')
IMPORTED = ROOT / '04-数据与复盘/平台导入/作品列表导出.xlsx'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
assert book.sheetnames == ['Sheet1'], book.sheetnames
rows = list(book.active.values)
headers = ['作品名称','发布时间','体裁','审核状态','播放量','完播率','5s完播率','封面点击率','2s跳出率','平均播放时长','点赞量','分享量','评论量','收藏量','主页访问量','粉丝增量']
assert list(rows[0]) == headers, rows[0]
records = []
for number, row in enumerate(rows[1:], 2):
    record = dict(zip(headers, row))
    assert record['作品名称'] and record['发布时间'], number
    published = record['发布时间']
    if isinstance(published, dt.datetime):
        published = published.strftime('%Y-%m-%d %H:%M:%S')
    else:
        dt.datetime.strptime(published, '%Y-%m-%d %H:%M:%S')
    record['发布时间'] = published
    record['source_row'] = number
    records.append(record)
book.close()
index_text = (ROOT / '00-内容资产索引/内容主索引.md').read_text(encoding='utf-8-sig')
match = re.search(r'<!-- dailywrite-delivery:v1 -->\s*```json\s*(.*?)\s*```', index_text, re.S)
assert match, 'Missing delivery metadata'
metadata = json.loads(match.group(1))
index_rows = []
for line in index_text.split('## 当前生产索引')[0].splitlines():
    if line.startswith('| DW-'):
        cells = [cell.strip() for cell in line.strip('|').split('|')]
        index_rows.append(dict(zip(['id','topic','status','platform','bodyPath','dataPath','review','reuse'], cells)))
old_data = []
for path in (ROOT / '04-数据与复盘/抖音数据').glob('*.md'):
    text = path.read_text(encoding='utf-8-sig')
    published = re.search(r'^发布时间[：:]\s*(.+)$', text, re.M)
    if published:
        matches = [r for r in records if r['发布时间'] == published.group(1).strip()]
        old_data.append({'file': path.name, 'published': published.group(1).strip(), 'sourceMatches': [{'title': r['作品名称'], 'visibility': r['审核状态'], 'row': r['source_row']} for r in matches]})
payload = {'source': str(SOURCE), 'sourceSha256': sha(SOURCE), 'importedSha256': sha(IMPORTED), 'records': records, 'indexRows': index_rows, 'metadata': metadata, 'oldData': old_data}
out = ROOT / 'output/douyin-data-refresh-20260929/source-inspection.json'
out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'recordCount': len(records), 'publicCount': sum(r['审核状态']=='公开' for r in records), 'sourceSha256': sha(SOURCE), 'importMatches': sha(SOURCE)==sha(IMPORTED), 'metadataKeys': list(metadata), 'indexRows': len(index_rows)}, ensure_ascii=False))
print(json.dumps(old_data, ensure_ascii=False))
for record in records:
    if record['审核状态'] == '公开':
        published_iso = record['发布时间'].replace(' ', 'T') + '+08:00'
        key_text = json.dumps([unicodedata.normalize('NFC', record['作品名称'].strip()), published_iso, record['体裁'].strip()], ensure_ascii=False, separators=(',', ':'))
        record_id = 'published-' + hashlib.sha256(key_text.encode('utf-8')).hexdigest()[:24]
        linked = metadata['platformLinks'].get(record_id)
        print(json.dumps({'row': record['source_row'], 'date': record['发布时间'], 'title': record['作品名称'].splitlines()[0], 'views': record['播放量'], 'recordId': record_id, 'contentId': linked['contentId'] if linked else None}, ensure_ascii=False))
for key, value in metadata.items():
    if isinstance(value, (list, dict)) and value:
        first = next(iter(value.items())) if isinstance(value, dict) else value[0]
        print(json.dumps({'metadataKey': key, 'sample': first}, ensure_ascii=False))
